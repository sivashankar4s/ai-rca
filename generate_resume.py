"""
Resume Generator - Sivashankar S
Produces:  Sivashankar_Resume.pdf   (fpdf2)
           Sivashankar_Resume.docx  (python-docx)
Run:  python generate_resume.py
"""

from fpdf import FPDF
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAME     = 'Sivashankar S'
TITLE    = 'Senior Backend Engineer  |  Agentic AI  |  Java  |  Python  |  AWS'
EMAIL    = 'sekarsivashankar4@gmail.com'
PHONE    = '+91-9688709534'
LINKEDIN = 'linkedin.com/in/sivashankar-s'
LOCATION = 'Chennai, India'

SUMMARY = (
    'Senior Backend Engineer with 9+ years of experience building scalable microservices and '
    'cloud-native systems. Expertise in Agentic AI, LLM integrations, and event-driven AWS '
    'architectures. Proven ability to deliver high-impact automation platforms and optimize '
    'large-scale data systems. Strong background in healthcare and fintech domains.'
)

SKILLS = [
    ('Languages',    'Java, Python, Node.js, JavaScript'),
    ('Frameworks',   'Spring Boot, Spring MVC, Flask'),
    ('Cloud / AWS',  'Lambda, S3, SQS, SNS, Glue, DynamoDB, EC2, CloudWatch'),
    ('AI / LLM',     'Agentic AI, LLM Integration, Mem0, Cursor, GitHub Copilot'),
    ('Databases',    'PostgreSQL, MySQL'),
    ('ORM',          'Hibernate, JPA'),
    ('Architecture', 'Microservices, CQRS, Event-driven, REST APIs'),
    ('DevOps',       'Docker, Jenkins, Maven, Gradle, Git'),
    ('Frontend',     'Angular, jQuery'),
    ('Other',        'Camunda BPM, Drools, SonarQube, JIRA'),
]

EXPERIENCE = [
    {
        'company': 'Ideas2IT Technology Services Pvt. Ltd.',
        'role':    'Senior Technical Analyst',
        'period':  'June 2016 - Present  (9+ years)',
        'points': [
            'Designed and developed scalable microservices and backend systems for enterprise clients.',
            'Led end-to-end feature delivery from architecture design to production deployment.',
            'Architected event-driven and CQRS-based solutions on AWS.',
            'Mentored junior engineers and drove code-quality improvements across teams.',
            'Improved system performance by 45% by rewriting SFTP file-fetching logic.',
            'Integrated Camunda BPM for workflow automation in enterprise finance workflows.',
            'Built FHIR-compliant migration tools for healthcare data across multi-tenant environments.',
        ],
    }
]

PROJECTS = [
    {
        'name':   'AI Job Automation Platform - Nurturebox (Bloom)',
        'domain': 'HR Tech / Recruitment',
        'points': [
            'Built Agentic AI platform for automated job discovery, personalisation, and auto-apply.',
            'Implemented LLM-based features: rephrase JDs and generate intelligent form answers.',
            'Developed browser-based auto-apply automation using Node.js.',
            'Integrated Mem0 to store and reuse responses for faster application completion.',
            'Designed scalable REST APIs for job search, tracking, and profile management.',
        ],
    },
    {
        'name':   'Healthcare Data Platform - Roche ONEDBM (Data Migration)',
        'domain': 'Healthcare',
        'points': [
            'Designed AWS serverless pipeline transforming large datasets into Parquet format.',
            'Built Lambda functions for raw-data transformation with SNS/SQS orchestration.',
            'Integrated AWS Glue for data cataloging and pipeline management.',
            'Implemented automated failure-handling and retry mechanisms.',
        ],
    },
    {
        'name':   'Navify Clinical Hub & Navify OncoHub',
        'domain': 'Healthcare (Roche)',
        'points': [
            'Developed multi-tenant microservices using Java Spring Boot for clinical decisions.',
            'Built patient journey system with AOP-driven chronological event management.',
            'Led application-level localization supporting multiple languages and regions.',
            'Utilized Drools rule engine for centralized business-rule management.',
        ],
    },
    {
        'name':   'OPORTUN - MMS (Loan Management)',
        'domain': 'Finance (FinTech)',
        'points': [
            'Integrated reconciliation and acknowledgment processes with Bank of America.',
            'Improved system performance by 45%+ through SFTP logic rewrite.',
            'Implemented Camunda BPM for template and approval workflow automation.',
            'Established CI/CD pipelines via Jenkins; managed IAM policies on AWS.',
        ],
    },
    {
        'name':   'DigiContent & CommerzEdge',
        'domain': 'E-Commerce',
        'points': [
            'Built product information management system for Walmart with JSON schema validation.',
            'Developed custom Ariba store using Reaction Commerce for seamless integration.',
            'Implemented automated build and deployment via Jenkins.',
        ],
    },
]

EDUCATION  = 'B.E. - Electrical & Electronics Engineering\nSri Venkateswara College of Engineering, 2016'
FUNCTIONAL = 'Camunda BPM  |  Reaction Commerce  |  JSON Schema Validators  |  FHIR  |  Drools'
ACHIEVEMENTS = [
    'Improved system performance by 45% through SFTP logic rewrite.',
    'Built AI-powered job automation, reducing manual effort significantly.',
    'Hands-on with Agentic AI tools: Cursor, GitHub Copilot.',
    'FHIR-compliant data migration across healthcare tenants.',
    'End-to-end Camunda BPM workflow integration in production.',
]


# ═══════════════════════════════════════════
#  PDF  (fpdf2)
# ═══════════════════════════════════════════

class ResumePDF(FPDF):
    def set_navy(self):   self.set_text_color(27,  42,  74)
    def set_blue(self):   self.set_text_color(46, 109, 164)
    def set_accent(self): self.set_text_color(74, 159, 212)
    def set_gray(self):   self.set_text_color(85,  85,  85)
    def set_mid(self):    self.set_text_color(136, 136, 136)
    def set_white(self):  self.set_text_color(255, 255, 255)

    def section_heading(self, text, w):
        self.set_font('Helvetica', 'B', 10)
        self.set_navy()
        self.cell(w, 6, text.upper(), ln=True)
        x = self.get_x(); y = self.get_y()
        self.set_draw_color(74, 159, 212)
        self.set_line_width(0.4)
        self.line(x, y, x + w, y)
        self.ln(2)
        self.set_draw_color(0,0,0); self.set_line_width(0.2)

    def header(self): pass
    def footer(self): pass


def make_pdf(filename='Sivashankar_Resume.pdf'):
    pdf = ResumePDF(orientation='P', unit='mm', format='A4')
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_margins(12, 8, 12)

    BODY_W = 186
    LEFT_W = 58
    RIGHT_W = BODY_W - LEFT_W - 3
    MARGIN = 12
    HDR_H = 28

    # Header
    pdf.set_fill_color(27, 42, 74)
    pdf.rect(MARGIN, 8, BODY_W, HDR_H, 'F')
    pdf.set_xy(MARGIN + 4, 10)
    pdf.set_font('Helvetica', 'B', 22)
    pdf.set_white()
    pdf.cell(0, 9, NAME, ln=True)
    pdf.set_x(MARGIN + 4)
    pdf.set_font('Helvetica', '', 9.5)
    pdf.set_accent()
    pdf.cell(0, 5, TITLE, ln=True)
    pdf.set_x(MARGIN + 4)
    pdf.set_font('Helvetica', '', 8)
    pdf.set_white()
    pdf.cell(0, 5, f'{EMAIL}   |   {PHONE}   |   {LINKEDIN}   |   {LOCATION}', ln=True)

    TOP_Y = 8 + HDR_H + 4
    X_LEFT = MARGIN
    pdf.set_xy(X_LEFT, TOP_Y)

    # LEFT column
    pdf.section_heading('Core Skills', LEFT_W)
    for cat, val in SKILLS:
        pdf.set_font('Helvetica', 'B', 7.5); pdf.set_navy()
        pdf.set_x(X_LEFT); pdf.cell(LEFT_W, 4, f'{cat}:', ln=True)
        pdf.set_font('Helvetica', '', 7.5); pdf.set_gray()
        pdf.set_x(X_LEFT); pdf.multi_cell(LEFT_W, 4, val, ln=True)
        pdf.ln(0.5)

    pdf.ln(2); pdf.set_x(X_LEFT)
    pdf.section_heading('Education', LEFT_W)
    pdf.set_font('Helvetica', '', 8); pdf.set_gray()
    for line in EDUCATION.split('\n'):
        pdf.set_x(X_LEFT); pdf.multi_cell(LEFT_W, 4.5, line, ln=True)

    pdf.ln(2); pdf.set_x(X_LEFT)
    pdf.section_heading('Achievements', LEFT_W)
    for ach in ACHIEVEMENTS:
        pdf.set_font('Helvetica', '', 7.8); pdf.set_gray()
        pdf.set_x(X_LEFT + 2); pdf.multi_cell(LEFT_W - 2, 4, f'* {ach}', ln=True)
        pdf.ln(0.5)

    pdf.ln(2); pdf.set_x(X_LEFT)
    pdf.section_heading('Functional', LEFT_W)
    pdf.set_font('Helvetica', '', 7.8); pdf.set_gray()
    for item in FUNCTIONAL.split('  |  '):
        pdf.set_x(X_LEFT); pdf.multi_cell(LEFT_W, 4, f'* {item.strip()}', ln=True)

    LEFT_END_Y = pdf.get_y()

    # RIGHT column
    X_RIGHT = MARGIN + LEFT_W + 3
    pdf.set_xy(X_RIGHT, TOP_Y)
    pdf.section_heading('Professional Summary', RIGHT_W)
    pdf.set_font('Helvetica', '', 9); pdf.set_gray()
    pdf.set_x(X_RIGHT); pdf.multi_cell(RIGHT_W, 5, SUMMARY, ln=True)

    pdf.ln(2); pdf.set_x(X_RIGHT)
    pdf.section_heading('Experience', RIGHT_W)
    for exp in EXPERIENCE:
        pdf.set_font('Helvetica', 'B', 9.5); pdf.set_navy()
        pdf.set_x(X_RIGHT); pdf.cell(RIGHT_W, 5, exp['company'], ln=True)
        pdf.set_font('Helvetica', 'I', 8); pdf.set_mid()
        pdf.set_x(X_RIGHT); pdf.cell(RIGHT_W, 4.5, f"{exp['role']}  |  {exp['period']}", ln=True)
        for pt in exp['points']:
            pdf.set_font('Helvetica', '', 8.5); pdf.set_gray()
            pdf.set_x(X_RIGHT + 3); pdf.multi_cell(RIGHT_W - 3, 4.5, f'*  {pt}', ln=True)
        pdf.ln(1)

    pdf.ln(1); pdf.set_x(X_RIGHT)
    pdf.section_heading('Key Projects', RIGHT_W)
    for proj in PROJECTS:
        pdf.set_font('Helvetica', 'B', 9); pdf.set_blue()
        pdf.set_x(X_RIGHT); pdf.cell(RIGHT_W, 5, proj['name'], ln=True)
        pdf.set_font('Helvetica', 'I', 7.5); pdf.set_mid()
        pdf.set_x(X_RIGHT); pdf.cell(RIGHT_W, 4, f"Domain: {proj['domain']}", ln=True)
        for pt in proj['points']:
            pdf.set_font('Helvetica', '', 8.2); pdf.set_gray()
            pdf.set_x(X_RIGHT + 3); pdf.multi_cell(RIGHT_W - 3, 4.5, f'*  {pt}', ln=True)
        pdf.ln(1.5)

    RIGHT_END_Y = pdf.get_y()

    # Column divider
    pdf.set_draw_color(200, 220, 238); pdf.set_line_width(0.3)
    pdf.line(MARGIN + LEFT_W + 1.5, TOP_Y, MARGIN + LEFT_W + 1.5, max(LEFT_END_Y, RIGHT_END_Y))

    pdf.output(filename)
    print(f'[OK] PDF created: {filename}')


# ═══════════════════════════════════════════
#  DOCX  (python-docx)
# ═══════════════════════════════════════════

def rgb(r, g, b): return RGBColor(r, g, b)
C_NAVY = rgb(27,42,74); C_BLUE = rgb(46,109,164); C_ACNT = rgb(74,159,212)
C_GRAY = rgb(85,85,85); C_MID = rgb(136,136,136); C_WHITE = rgb(255,255,255)

def _set_cell_bg(cell, hex6):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr()
    shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hex6)
    tcPr.append(shd)

def _bottom_border(para, color='4A9FD4'):
    pPr=para._p.get_or_add_pPr(); pBdr=OxmlElement('w:pBdr')
    b=OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'6')
    b.set(qn('w:space'),'1'); b.set(qn('w:color'),color)
    pBdr.append(b); pPr.append(pBdr)

def _left_border_cell(cell, color='CCDDEE'):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcB=OxmlElement('w:tcBorders')
    l=OxmlElement('w:left')
    l.set(qn('w:val'),'single'); l.set(qn('w:sz'),'4')
    l.set(qn('w:space'),'0'); l.set(qn('w:color'),color)
    tcB.append(l); tcPr.append(tcB)

def _r(para, text, bold=False, italic=False, size=None, color=None):
    run=para.add_run(text); run.bold=bold; run.italic=italic; run.font.name='Calibri'
    if size:  run.font.size=Pt(size)
    if color: run.font.color.rgb=color
    return run

def _sec(cell, text):
    p=cell.add_paragraph()
    p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(2)
    _r(p,text.upper(),bold=True,size=10,color=C_NAVY); _bottom_border(p); return p

def _bullet(cell, text, size=8.5, color=C_GRAY):
    p=cell.add_paragraph()
    p.paragraph_format.left_indent=Cm(0.3)
    p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
    _r(p,f'*  {text}',size=size,color=color); return p

def _clear(cell):
    for p in list(cell.paragraphs): p._element.getparent().remove(p._element)

def _build_left(cell):
    _clear(cell)
    p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(2)
    _r(p,'CORE SKILLS',bold=True,size=10,color=C_NAVY); _bottom_border(p)
    for cat,val in SKILLS:
        p2=cell.add_paragraph(); p2.paragraph_format.space_before=Pt(2); p2.paragraph_format.space_after=Pt(1)
        _r(p2,f'{cat}: ',bold=True,size=8,color=C_NAVY); _r(p2,val,size=7.8,color=C_GRAY)
    _sec(cell,'Education')
    p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    for line in EDUCATION.split('\n'): _r(p,line+'\n',size=8.5,color=C_GRAY)
    _sec(cell,'Achievements')
    for ach in ACHIEVEMENTS:
        pb=cell.add_paragraph(); pb.paragraph_format.left_indent=Cm(0.2)
        pb.paragraph_format.space_before=Pt(1); pb.paragraph_format.space_after=Pt(1)
        _r(pb,f'* {ach}',size=7.8,color=C_GRAY)
    _sec(cell,'Functional')
    for item in FUNCTIONAL.split('  |  '):
        pb=cell.add_paragraph(); pb.paragraph_format.left_indent=Cm(0.2)
        pb.paragraph_format.space_before=Pt(1); pb.paragraph_format.space_after=Pt(1)
        _r(pb,f'* {item.strip()}',size=8,color=C_GRAY)

def _build_right(cell):
    _clear(cell)
    _sec(cell,'Professional Summary')
    p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(4)
    _r(p,SUMMARY,size=9,color=C_GRAY)
    _sec(cell,'Experience')
    for exp in EXPERIENCE:
        p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(3)
        _r(p,exp['company'],bold=True,size=9.5,color=C_NAVY)
        p2=cell.add_paragraph(); p2.paragraph_format.space_after=Pt(2)
        _r(p2,f"{exp['role']}  |  {exp['period']}",italic=True,size=8,color=C_MID)
        for pt in exp['points']: _bullet(cell,pt,size=8.5,color=C_GRAY)
    _sec(cell,'Key Projects')
    for proj in PROJECTS:
        p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(4)
        _r(p,proj['name'],bold=True,size=9,color=C_BLUE)
        p2=cell.add_paragraph(); p2.paragraph_format.space_after=Pt(1)
        _r(p2,f"Domain: {proj['domain']}",italic=True,size=7.5,color=C_MID)
        for pt in proj['points']: _bullet(cell,pt,size=8.2,color=C_GRAY)

def make_docx(filename='Sivashankar_Resume.docx'):
    doc=Document()
    for sec in doc.sections:
        sec.top_margin=Cm(0.8); sec.bottom_margin=Cm(1.2)
        sec.left_margin=Cm(1.3); sec.right_margin=Cm(1.3)
    # Header
    ht=doc.add_table(rows=1,cols=1); ht.alignment=WD_TABLE_ALIGNMENT.CENTER
    hc=ht.cell(0,0); _set_cell_bg(hc,'1B2A4A')
    pn=hc.paragraphs[0]; pn.paragraph_format.space_before=Pt(8); pn.paragraph_format.space_after=Pt(2)
    pn.alignment=WD_ALIGN_PARAGRAPH.LEFT; _r(pn,NAME,bold=True,size=22,color=C_WHITE)
    pt=hc.add_paragraph(); pt.alignment=WD_ALIGN_PARAGRAPH.LEFT; pt.paragraph_format.space_after=Pt(2)
    _r(pt,TITLE,size=9.5,color=C_ACNT)
    pc=hc.add_paragraph(); pc.alignment=WD_ALIGN_PARAGRAPH.LEFT; pc.paragraph_format.space_after=Pt(8)
    _r(pc,f'{EMAIL}   |   {PHONE}   |   {LINKEDIN}   |   {LOCATION}',size=8.5,color=C_WHITE)
    sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(2)
    # Two-col body
    bt=doc.add_table(rows=1,cols=2); bt.alignment=WD_TABLE_ALIGNMENT.CENTER
    bt.columns[0].width=Cm(5.5); bt.columns[1].width=Cm(11.5)
    lc=bt.cell(0,0); rc=bt.cell(0,1)
    lc.vertical_alignment=WD_ALIGN_VERTICAL.TOP; rc.vertical_alignment=WD_ALIGN_VERTICAL.TOP
    _build_left(lc); _build_right(rc); _left_border_cell(rc)
    doc.save(filename)
    print(f'[OK] DOCX created: {filename}')

if __name__ == '__main__':
    make_pdf('Sivashankar_Resume.pdf')
    make_docx('Sivashankar_Resume.docx')
    print('\nDone!')
